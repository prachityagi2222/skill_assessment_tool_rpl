import os
import openpyxl
import pypdfium2 as pdfium
import json
import re
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = str(BASE_DIR / "data")

# Check NQR Excel for specific codes and statistics
excel_path = os.path.join(DATA_DIR, "NQR_Qualifications_Summary.xlsx.xlsx")
wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
ws = wb['Worksheet']
excel_rows = list(ws.iter_rows(values_only=True))
headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(excel_rows[2])]

print(f"Excel Headers ({len(headers)}): {headers}")
print(f"Total rows in Excel: {len(excel_rows)}")

# Find MEPSC code in Excel
mepsc_rows = []
es_rows = []
for idx, r in enumerate(excel_rows[3:], start=1):
    r_dict = dict(zip(headers, r))
    code = str(r_dict.get('Code', '') or '')
    title = str(r_dict.get('Title', '') or '')
    if '06022' in code or 'Designer- Assessment' in title or 'MEPSC' in code:
        mepsc_rows.append((idx, r_dict))
    if 'Employability' in title or 'Employability' in str(r_dict.get('Description', '')):
        es_rows.append((idx, r_dict))

print(f"Matches for MEPSC / Assessment Designer in Excel: {len(mepsc_rows)}")
for idx, r in mepsc_rows[:5]:
    print(f"Row {idx}: Code={r.get('Code')}, Title={r.get('Title')}, Level={r.get('Level')}")

print(f"Matches for Employability in Excel: {len(es_rows)}")
for idx, r in es_rows[:5]:
    print(f"Row {idx}: Code={r.get('Code')}, Title={r.get('Title')}, Level={r.get('Level')}")

# Now inspect Employability Skills PDFs
pdf_names = [
    "Employability_Skills_30_Hours.pdf.pdf",
    "Employability_Skills_60_Hours.pdf.pdf",
    "Employability_Skills_90_Hours.pdf.pdf",
    "Employability_Skills_120_Hours.pdf.pdf"
]

es_extracted = {}
for p_name in pdf_names:
    fpath = os.path.join(DATA_DIR, p_name)
    doc = pdfium.PdfDocument(fpath)
    full_text = []
    for i, p in enumerate(doc):
        t = p.get_textpage().get_text_range()
        full_text.append({"page": i+1, "text": t})
    
    combined = "\n".join([p["text"] for p in full_text])
    duration = re.search(r'\((\d+)\s*hours?\)', combined, re.IGNORECASE)
    duration_val = duration.group(1) if duration else "Unknown"
    
    es_extracted[p_name] = {
        "file": p_name,
        "pages_count": len(doc),
        "duration_hours": duration_val,
        "sample_start": full_text[0]["text"][:600] if full_text else "",
        "full_pages": full_text
    }

with open("es_extracted_raw.json", "w", encoding="utf-8") as f:
    json.dump(es_extracted, f, indent=2, ensure_ascii=False)

print("Saved es_extracted_raw.json successfully!")
